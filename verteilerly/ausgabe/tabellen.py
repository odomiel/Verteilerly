"""Excel und CSV: die Belegungsdaten je Ader.

Die Tabellenformate geben nicht das Blatt wieder, sondern die Daten - eine
Zeile je Ader. Damit laesst sich filtern, sortieren und weiterrechnen, was
auf einem Blatt nur zu betrachten waere.
"""

from __future__ import annotations

import csv
import pathlib

from ..blattvorlagen import STANDARD, Blattvorlage
from ..layout import zellentext
from ..modell import (
    BRIEFKOPF_BESCHRIFTUNG,
    Aderadresse,
    Projekt,
)
from . import AusgabeFehler

SPALTEN = (
    "Bucht",
    "Leiste",
    "Stift",
    "Ader",
    "Adresse",
    "Farbe",
    "Bezeichnung",
    "Rangierung",
    "Bemerkung",
    "Kabel",
    "Kabeltyp",
)
KABELSPALTEN = (
    "Beschriftung",
    "Kabeltyp",
    "DA gesamt",
    "DA aufgelegt",
    "gekürzt",
    "Abschnitte",
)
TRENNER = ";"  # deutsches Excel oeffnet CSV so ohne Rueckfrage


def belegungszeilen(projekt: Projekt) -> list[list]:
    """Eine Zeile je Ader, in der Reihenfolge des Verteilers.

    Bucht, Leiste und Stift bleiben Zahlen, damit in Excel richtig
    sortiert wird und nicht 10 vor 2 steht.
    """
    v = projekt.verteiler
    zeilen: list[list] = []
    for bucht in v.buchten:
        for leiste in bucht.leisten:
            for stift in leiste.stifte:
                kabel = v.kabel_an(leiste.nummer, stift.nummer)
                for seite in ("a", "b"):
                    ader = stift.ader(seite)
                    adresse = Aderadresse(leiste.nummer, stift.nummer, seite)
                    zeilen.append(
                        [
                            bucht.nummer,
                            leiste.nummer,
                            stift.nummer,
                            seite,
                            adresse.anzeige(projekt.schreibweise),
                            ader.farbe,
                            zellentext(ader.bezeichnung, ader.teilung("bezeichnung")),
                            v.rangierfeld(adresse, projekt.schreibweise),
                            zellentext(ader.bemerkung, ader.teilung("bemerkung")),
                            kabel.beschriftung if kabel else "",
                            kabel.kabeltyp if kabel else "",
                        ]
                    )
    return zeilen


def kabelzeilen(projekt: Projekt) -> list[list]:
    zeilen = []
    for kabel in projekt.verteiler.kabel:
        abschnitte = ", ".join(
            f"Leiste {a.leiste} DA {a.start_da}–{a.end_da}" for a in kabel.abschnitte
        )
        zeilen.append(
            [
                kabel.beschriftung,
                kabel.kabeltyp,
                kabel.da_gesamt or kabel.da_aufgelegt,
                kabel.da_aufgelegt,
                "ja" if kabel.gekuerzt else "nein",
                abschnitte,
            ]
        )
    return zeilen


def projektdatenzeilen(projekt: Projekt) -> list[list]:
    """Der Briefkopf als Liste - Beschriftung und Wert.

    Dieselben Zeilen fuellen das Blatt der Mappe und den Fussblock der
    CSV: zwei Auslesungen liefen auseinander.
    """
    zeilen = [
        ["Verteiler", projekt.verteiler.name],
        ["Verteilernummer", projekt.verteiler.nummer],
    ]
    zeilen += [
        [BRIEFKOPF_BESCHRIFTUNG.get(name, name), wert]
        for name, wert in projekt.briefkopf.items()
        if wert
    ]
    zeilen += [[name, wert] for name, wert in projekt.eigene_felder.items()]
    return [zeile for zeile in zeilen if str(zeile[1]).strip()]


def csv_bloecke(projekt: Projekt) -> list[tuple[str, tuple, list]]:
    """Was in der CSV steht: je Block Titel, Ueberschriften und Zeilen.

    Die Belegung steht oben und ohne Titel - die erste Zeile bleibt die
    Kopfzeile der Tabelle, wie es jedes einlesende Programm erwartet.
    Was auf dem Blatt im Schriftfeld steht, kommt darunter, so wie es
    dort am Blattfuss steht.
    """
    bloecke = [("", SPALTEN, belegungszeilen(projekt))]
    kabel = kabelzeilen(projekt)
    if kabel:
        bloecke.append(("Kabel je Leiste", KABELSPALTEN, kabel))
    stammdaten = projektdatenzeilen(projekt)
    if stammdaten:
        bloecke.append(("Projektdaten", ("Feld", "Wert"), stammdaten))
    return bloecke


def csv_schreiben(projekt: Projekt, ziel: pathlib.Path) -> pathlib.Path:
    """Die Belegung, darunter Kabel und Projektdaten.

    Eine CSV traegt nur eine Tabelle - also stehen die weiteren als
    eigene Bloecke darunter, durch eine Leerzeile getrennt und mit
    eigener Kopfzeile. Oben bleibt die Belegung: die erste Zeile ist
    ihre Kopfzeile, damit eingelesen werden kann wie bisher.
    """
    ziel = pathlib.Path(ziel)
    try:
        # BOM, damit Excel die Umlaute nicht zerlegt.
        with ziel.open("w", newline="", encoding="utf-8-sig") as datei:
            schreiber = csv.writer(datei, delimiter=TRENNER)
            for i, (titel, ueberschriften, zeilen) in enumerate(csv_bloecke(projekt)):
                if i:
                    schreiber.writerow([])
                if titel:
                    schreiber.writerow([titel])
                schreiber.writerow(list(ueberschriften))
                # Eine Anschrift steht mehrzeilig in einer Zelle; die CSV
                # gibt sie in einer Zeile wieder - sonst zerfiele die
                # Tabelle an dieser Stelle in drei Zeilen.
                schreiber.writerows(
                    [
                        [
                            str(wert).replace("\n", " · ")
                            if isinstance(wert, str) else wert
                            for wert in zeile
                        ]
                        for zeile in zeilen
                    ]
                )
    except OSError as fehler:
        raise AusgabeFehler(str(fehler)) from fehler
    return ziel


def xlsx_schreiben(
    projekt: Projekt, ziel: pathlib.Path, vorlage: Blattvorlage = STANDARD
) -> pathlib.Path:
    """Die Blaetter im Blattbild, dahinter die Daten.

    Vorn steht, was auch gedruckt wird - Deckblatt, Verteilerbelegung und
    je Leiste ein Blatt, gesetzt nach der gewaehlten Blattvorlage. Dahinter
    liegen dieselben Daten als Tabelle, zum Filtern und Weiterrechnen.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError as fehler:  # pragma: no cover - haengt an der Umgebung
        raise AusgabeFehler("openpyxl ist nicht vorhanden.") from fehler

    from .mappe import blaetter_eintragen

    ziel = pathlib.Path(ziel)
    mappe = Workbook()
    blaetter_eintragen(mappe, projekt, vorlage)

    belegung = mappe.create_sheet("Daten je Ader")
    _blatt_fuellen(belegung, SPALTEN, belegungszeilen(projekt), Font, Alignment,
                   get_column_letter)
    belegung.freeze_panes = "A2"
    belegung.auto_filter.ref = belegung.dimensions

    kabel = mappe.create_sheet("Kabel je Leiste")
    _blatt_fuellen(kabel, KABELSPALTEN, kabelzeilen(projekt), Font, Alignment,
                   get_column_letter)

    stammdaten = mappe.create_sheet("Projektdaten")
    _blatt_fuellen(stammdaten, ("Feld", "Wert"), projektdatenzeilen(projekt),
                   Font, Alignment, get_column_letter)

    try:
        mappe.save(ziel)
    except OSError as fehler:
        raise AusgabeFehler(str(fehler)) from fehler
    return ziel


def _blatt_fuellen(blatt, ueberschriften, zeilen, Font, Alignment, spaltenname) -> None:
    blatt.append(list(ueberschriften))
    for zelle in blatt[1]:
        zelle.font = Font(bold=True)
        zelle.alignment = Alignment(vertical="center")
    for zeile in zeilen:
        blatt.append(zeile)
        # Eine Anschrift steht mehrzeilig in der Zelle - ohne Umbruch
        # zeigte Excel nur die erste Zeile.
        for zelle in blatt[blatt.max_row]:
            if isinstance(zelle.value, str) and "\n" in zelle.value:
                zelle.alignment = Alignment(vertical="top", wrap_text=True)

    for nummer, name in enumerate(ueberschriften, 1):
        laengste = max(
            [len(str(name))] + [len(str(z[nummer - 1])) for z in zeilen] or [0]
        )
        blatt.column_dimensions[spaltenname(nummer)].width = min(max(laengste + 2, 9), 46)
