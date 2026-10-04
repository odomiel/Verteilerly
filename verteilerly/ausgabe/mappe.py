"""Excel im Blattbild: dieselben Blaetter, nur als Tabelle.

Excel ist ein Raster, das Blatt eine Zeichnung - genau nachbauen laesst
sich ein PDF hier nicht. Das Leistenblatt ist aber selbst eine Tabelle:
Spalten, Kopfzeile, zwei Zeilen je Stift. Genau das wird hier gesetzt,
mit den Massen und Schriftgroessen der gewaehlten Blattvorlage, und zwar
so, dass ein Ausdruck aus Excel dem PDF nahekommt.

Was Excel nicht nachbilden kann, ist die freie Lage in Millimetern. Der
Kabelblock wird darum ueber seine Zeilen verbunden statt frei gezeichnet,
und das Schriftfeld steht unter der Tabelle statt am Blattfuss.
"""

from __future__ import annotations

from ..blattvorlagen import STANDARD, Blattvorlage
from ..layout import (
    BESCHRIFTUNGSANTEIL,
    FUELLUNG_KABEL,
    FUELLUNG_KOPF,
    LINIE,
    LINIE_FEIN,
    schriftfeldanteile,
    beispielschriftfeld,
    belegungsgruppen,
    belegungsmasse,
    buchtnamen,
    deckblattzeilen,
    kabelfeld,
    kopfgruppen,
    STRICH,
    schriftfeldspalten,
    schriftfeldwerte,
    schriftfeldzeilen,
    umgebrochene_zeilen,
    zeilenstile,
    zeilenhoehe_leiste,
    zeilenteilungen,
    zeilenwerte,
    zellentext,
)
from ..modell import LEERER_STIL, UNGETEILT, Projekt
from ..sprachen import u

SCHRIFTART = "Arial"          # was Helvetica im PDF am naechsten kommt
PT_JE_MM = 72.0 / 25.4
PX_JE_MM = 96.0 / 25.4
# Excel misst Spalten in Zeichenbreiten der Standardschrift: ein Zeichen
# sind rund sieben Bildpunkte, dazu kommt der Zellrand.
PX_JE_ZEICHEN = 7.0
ZELLRAND_PX = 5.0


def spaltenbreite(mm: float) -> float:
    """Millimeter in Excels Spaltenmass."""
    return max(1.0, round((mm * PX_JE_MM - ZELLRAND_PX) / PX_JE_ZEICHEN, 2))


def zeilenhoehe(mm: float) -> float:
    """Millimeter in Punkt - so misst Excel Zeilenhoehen."""
    return round(mm * PT_JE_MM, 2)


# Wie hoch eine Zeile ist, die niemand gesetzt hat - Excels Vorgabe.
VORGABEHOEHE_PT = 15.0
# Der Seitenrand aus _seiteneinrichtung, in Millimetern.
SEITENRAND_MM = 0.5 * 25.4
# Excel nimmt keine Zeile hoeher als 409,5 Punkt.
HOECHSTE_ZEILE_PT = 409.0


def _hex(farbe: str) -> str:
    return "FF" + farbe.lstrip("#").upper()


class _Werkzeug:
    """Die Bausteine von openpyxl, einmal geholt und benannt."""

    def __init__(self) -> None:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        self.Alignment = Alignment
        self.Border = Border
        self.Font = Font
        self.Fuellung = PatternFill
        self.Side = Side
        self.spaltenname = get_column_letter

    def strich(self, farbe: str, kraeftig: bool = False):
        return self.Side(
            style="thin" if not kraeftig else "medium", color=_hex(farbe)
        )

    def rahmen(self, **seiten):
        return self.Border(**seiten)

    def schrift(self, groesse: float, fett: bool = False, farbe: str = "#000000",
                kursiv: bool = False):
        return self.Font(
            name=SCHRIFTART, size=groesse, bold=fett, italic=kursiv,
            color=_hex(farbe),
        )

    def lage(self, ausrichtung: str = "links", gedreht: int = 0, umbruch: bool = False):
        waagerecht = {"links": "left", "mitte": "center", "rechts": "right"}
        return self.Alignment(
            horizontal=waagerecht.get(ausrichtung, "left"),
            vertical="center",
            text_rotation=gedreht,
            wrap_text=umbruch,
        )


def blaetter_eintragen(
    mappe, projekt: Projekt, vorlage: Blattvorlage = STANDARD
) -> list:
    """Legt Deckblatt, Verteilerbelegung und die Leistenblaetter an."""
    zeug = _Werkzeug()
    blaetter = [_deckblatt(mappe, projekt, vorlage, zeug)]
    blaetter.extend(_belegungen(mappe, projekt, vorlage, zeug))
    for leiste in projekt.verteiler.alle_leisten:
        blaetter.append(_leistenblatt(mappe, projekt, leiste.nummer, vorlage, zeug))
    return blaetter


def _seiteneinrichtung(blatt, quer: bool = False) -> None:
    """A4, damit ein Ausdruck aus Excel wie das PDF aussieht.

    Ein Blatt ist ein Blatt: Breite **und** Hoehe werden auf eine Seite
    gebracht. Die Zeilen tragen die Masse des Blattmodells, also
    fuellt die Seite sich von selbst; was dann noch ein paar Millimeter
    uebersteht - eine Ueberschrift, eine Leerzeile -, faengt Excel mit
    seiner Skalierung ab, statt eine zweite Seite anzufangen.
    """
    blatt.page_setup.paperSize = blatt.PAPERSIZE_A4
    blatt.page_setup.orientation = "landscape" if quer else "portrait"
    blatt.page_setup.fitToWidth = 1
    blatt.page_setup.fitToHeight = 1
    blatt.sheet_properties.pageSetUpPr.fitToPage = True
    blatt.page_margins.left = blatt.page_margins.right = 0.5
    blatt.page_margins.top = blatt.page_margins.bottom = 0.5


# -- Deckblatt --------------------------------------------------------------


def _deckblatt(mappe, projekt: Projekt, vorlage: Blattvorlage, zeug) -> object:
    blatt = mappe.active
    blatt.title = "Deckblatt"
    v = projekt.verteiler

    blatt.column_dimensions["A"].width = 26
    blatt.column_dimensions["B"].width = 52

    blatt["A1"] = projekt.anzeigewert("dokumenttitel", "Verteilerdokumentation")
    blatt["A1"].font = zeug.schrift(vorlage.schrift_gross, True)
    blatt.row_dimensions[1].height = zeilenhoehe(12)

    blatt["A3"] = v.name
    blatt["A3"].font = zeug.schrift(vorlage.schrift_mittel, True)
    blatt["B3"] = v.nummer
    blatt["B3"].font = zeug.schrift(vorlage.schrift_mittel)

    # Dieselbe Liste wie auf dem gezeichneten Deckblatt:
    # Aufbau, Kennzahlen, Anschriften, Datum und Name.
    zeile = 5
    for name, wert in deckblattzeilen(projekt, vorlage):
        if name:
            blatt.cell(row=zeile, column=1, value=name).font = zeug.schrift(
                vorlage.schrift_normal, False, "#555555"
            )
        feld = blatt.cell(row=zeile, column=2, value=wert)
        feld.font = zeug.schrift(vorlage.schrift_mittel)
        blatt.row_dimensions[zeile].height = zeilenhoehe(7.6)
        zeile += 1

    _seiteneinrichtung(blatt)
    return blatt


# -- Verteilerbelegung ------------------------------------------------------


def _belegungsspalte(
    da_anzahl: int, start_da: int, anzahl_da: int, umgekehrt: bool
) -> int:
    """Die erste Spalte eines DA-Blocks in der Leiste - ab 1 gezaehlt.

    Dasselbe wie ``layout._belegungsplatz``, nur in Zellen statt in
    Millimetern: in der Normalform steht DA 1 links, umgekehrt rechts.
    """
    vorn = start_da - 1
    if umgekehrt:
        vorn = da_anzahl - vorn - anzahl_da
    return vorn + 1


def _belegungen(mappe, projekt: Projekt, vorlage: Blattvorlage, zeug) -> list:
    """Je Blatt der Zeichnung eines in der Mappe."""
    gruppen = belegungsgruppen(projekt, vorlage)
    blaetter = []
    for gruppe in gruppen:
        name = "Verteilerbelegung" if len(gruppen) == 1 else buchtnamen(gruppe)
        blaetter.append(_belegung(mappe, projekt, gruppe, name, vorlage, zeug))
    return blaetter


def _belegung(
    mappe, projekt: Projekt, buchten: list, name: str, vorlage: Blattvorlage, zeug
) -> object:
    """Jede Leiste eine Zeile aus DA-Zellen, Kabel als verbundener Block.

    Die Masse kommen aus dem Blattmodell: Zellen, Beschriftung und Luft
    so gross wie dort, damit die Seite ebenso ausgefuellt ist.
    Die Buchten stehen nebeneinander, nicht untereinander - wie im PDF.
    """
    blatt = mappe.create_sheet(name)
    v = projekt.verteiler
    masse = belegungsmasse(buchten, vorlage)
    # Eine Leerzeile zwischen den Leisten traegt die Luft des Blattes.
    mit_luft = masse.luft_leisten > 0.5

    ueberschrift = blatt.cell(row=1, column=1, value=name)
    ueberschrift.font = zeug.schrift(vorlage.schrift_mittel, True)
    blatt.row_dimensions[1].height = zeilenhoehe(8)
    blatt.row_dimensions[2].height = zeilenhoehe(4)

    erste = 3
    spalte = 1
    letzte_spalte = 1
    letzte_zeile = erste
    for bucht in buchten:
        breiteste = max((l.da_anzahl for l in bucht.leisten), default=0)
        blatt.column_dimensions[zeug.spaltenname(spalte)].width = spaltenbreite(
            masse.label_breite
        )
        for i in range(breiteste):
            blatt.column_dimensions[
                zeug.spaltenname(spalte + 1 + i)
            ].width = spaltenbreite(masse.zelle_breite)

        kopf = blatt.cell(row=erste, column=spalte, value=buchtnamen([bucht]))
        kopf.font = zeug.schrift(vorlage.schrift_normal, True)

        zeile = erste + 1
        for leiste in bucht.leisten:
            beschriftung = blatt.cell(
                row=zeile, column=spalte,
                value=u("Leiste {0}").format(projekt.leistenname(leiste.nummer)),
            )
            beschriftung.font = zeug.schrift(masse.schrift, False, "#555555")
            beschriftung.alignment = zeug.lage("links")
            blatt.row_dimensions[zeile].height = zeilenhoehe(masse.zelle_hoehe)

            # Wie auf dem Blatt: kein Raster, nur der Kasten um die
            # Leiste - und DA-Nummern nur, wenn die Vorlage sie führt.
            for da in range(1, leiste.da_anzahl + 1):
                feld = blatt.cell(
                    row=zeile,
                    column=spalte + _belegungsspalte(
                        leiste.da_anzahl, da, 1, projekt.belegung_umgekehrt
                    ),
                    value=str(da) if vorlage.da_nummern else None,
                )
                if vorlage.da_nummern:
                    feld.font = zeug.schrift(
                        max(vorlage.schrift_klein - 1, 5), False, "#555555"
                    )
                    feld.alignment = zeug.lage("mitte")
                # Der Kasten umschliesst die Leiste - links ist links,
                # gleich welche Doppelader dort steht.
                stelle = feld.column - spalte
                feld.border = zeug.rahmen(
                    left=zeug.strich(LINIE_FEIN) if vorlage.da_nummern
                    else (zeug.strich(LINIE) if stelle == 1 else None),
                    right=zeug.strich(LINIE_FEIN) if vorlage.da_nummern
                    else (
                        zeug.strich(LINIE)
                        if stelle == leiste.da_anzahl
                        else None
                    ),
                    top=zeug.strich(LINIE), bottom=zeug.strich(LINIE),
                )

            for kabel in v.kabel:
                for abschnitt in kabel.abschnitte:
                    if abschnitt.leiste != leiste.nummer:
                        continue
                    von = spalte + _belegungsspalte(
                        leiste.da_anzahl, abschnitt.start_da,
                        abschnitt.anzahl_da, projekt.belegung_umgekehrt,
                    )
                    bis = von + abschnitt.anzahl_da - 1
                    blatt.merge_cells(
                        start_row=zeile, start_column=von,
                        end_row=zeile, end_column=bis,
                    )
                    # Wie auf dem Blatt: mitten im Block, so gross wie er
                    # es traegt, und umgebrochen statt abgeschnitten. Die
                    # Zeilen rechnet das Blattmodell - Excel kennt die
                    # Breite des verbundenen Blocks nicht und schriebe
                    # sonst unverkleinert darueber hinaus.
                    feld = kabelfeld(
                        0.0, 0.0,
                        abschnitt.anzahl_da * masse.zelle_breite,
                        masse.zelle_hoehe,
                        vorlage,
                    )
                    groesse, zeilen = umgebrochene_zeilen(
                        kabel.beschriftung, feld.breite, feld.hoehe,
                        feld.groesse, feld.einzeilig_ab,
                    )
                    feld = blatt.cell(
                        row=zeile, column=von, value="\n".join(zeilen)
                    )
                    feld.font = zeug.schrift(groesse, True)
                    feld.alignment = zeug.lage("mitte", umbruch=True)
                    feld.fill = zeug.Fuellung(
                        "solid", start_color=_hex(FUELLUNG_KABEL)
                    )
            zeile += 1
            if mit_luft and leiste is not bucht.leisten[-1]:
                blatt.row_dimensions[zeile].height = zeilenhoehe(masse.luft_leisten)
                zeile += 1
        letzte_zeile = max(letzte_zeile, zeile)
        letzte_spalte = spalte + breiteste
        # Zwischen den Buchten steht eine Spalte Luft.
        spalte = letzte_spalte + 1
        blatt.column_dimensions[zeug.spaltenname(spalte)].width = spaltenbreite(
            max(masse.luft_buchten, 2.0)
        )
        spalte += 1

    # Auch dieses Blatt traegt das Schriftfeld - im PDF steht es
    # am Blattfuss, hier unter der Zeichnung.
    fuss = _bis_zum_blattfuss(
        blatt, vorlage, letzte_zeile + 1, quer=not masse.hochkant
    )
    _schriftfeld(
        blatt, projekt, None, vorlage, zeug, fuss,
        max(letzte_spalte, 4), "belegung",
    )

    # Wie auf dem gezeichneten Blatt: eine Bucht steht hochkant, ab zwei
    # quer.
    _seiteneinrichtung(blatt, quer=not masse.hochkant)
    return blatt


# -- Leistenblatt -----------------------------------------------------------


def _leistenblatt(
    mappe, projekt: Projekt, leiste_nr: int, vorlage: Blattvorlage, zeug
) -> object:
    v = projekt.verteiler
    leiste = v.leiste(leiste_nr)
    blatt = mappe.create_sheet(_blatttitel(projekt, leiste_nr, "leisten"))
    spalten = vorlage.spalten
    breiten = vorlage.breiten(210.0 - 2 * vorlage.rand)

    for i, breite in enumerate(breiten, 1):
        blatt.column_dimensions[zeug.spaltenname(i)].width = spaltenbreite(breite)

    # Kein Blattkopf über der Tabelle - wie auf dem gezeichneten Blatt:
    # welche Leiste das Blatt zeigt, sagt das Schriftfeld darunter, und
    # das Register trägt den Namen. Der Platz gehört der Tabelle.
    KOPFZEILE = 1
    for i, spalte in enumerate(spalten, 1):
        feld = blatt.cell(row=KOPFZEILE, column=i, value=spalte.ueberschrift)
        feld.font = zeug.schrift(vorlage.schrift_klein, True)
        feld.alignment = zeug.lage(spalte.ausrichtung)
        feld.fill = zeug.Fuellung("solid", start_color=_hex(FUELLUNG_KOPF))
        feld.border = zeug.rahmen(
            top=zeug.strich(LINIE, True), bottom=zeug.strich(LINIE, True),
            left=zeug.strich(LINIE_FEIN), right=zeug.strich(LINIE_FEIN),
        )
    # „Stift“ steht über der Nummer und über a und b - wie auf dem Blatt.
    for erste, anzahl in kopfgruppen(spalten):
        if anzahl > 1:
            blatt.merge_cells(
                start_row=KOPFZEILE, start_column=erste + 1,
                end_row=KOPFZEILE, end_column=erste + anzahl,
            )
    blatt.row_dimensions[KOPFZEILE].height = zeilenhoehe(vorlage.kopfzeile_hoehe)

    erste_zeile = KOPFZEILE + 1
    zeile = erste_zeile
    for stift in leiste.stifte:
        for seite in ("a", "b"):
            werte = zeilenwerte(projekt, leiste_nr, stift.nummer, seite)
            stile = zeilenstile(projekt, leiste_nr, stift.nummer, seite)
            teilungen = zeilenteilungen(projekt, leiste_nr, stift.nummer, seite)
            for i, spalte in enumerate(spalten, 1):
                # Eine geteilte Zelle gibt es im Raster nicht: beide
                # Felder stehen hier hintereinander.
                text = zellentext(
                    werte.get(spalte.feld, ""),
                    teilungen.get(spalte.feld, UNGETEILT),
                )
                feld = blatt.cell(row=zeile, column=i, value=text)
                # Derselbe Zellenstil wie auf dem Blatt.
                stil = stile.get(spalte.feld, LEERER_STIL)
                feld.font = zeug.schrift(
                    vorlage.schrift_normal, stil.fett, kursiv=stil.kursiv
                )
                feld.alignment = zeug.lage(stil.ausrichtung or spalte.ausrichtung)
                # Unter der b-Zeile trennt ein kraeftiger Strich die Stifte.
                feld.border = zeug.rahmen(
                    left=zeug.strich(LINIE_FEIN), right=zeug.strich(LINIE_FEIN),
                    bottom=zeug.strich(LINIE, seite == "b"),
                )
            # Dieselbe Zeilenhoehe wie auf dem Blatt: die Tabelle nutzt
            # die Seite aus, statt als Streifen oben zu stehen.
            blatt.row_dimensions[zeile].height = zeilenhoehe(
                zeilenhoehe_leiste(leiste.da_anzahl, vorlage)
            )
            zeile += 1

    _kabelbloecke(blatt, projekt, leiste_nr, vorlage, zeug, erste_zeile)
    fuss = _bis_zum_blattfuss(blatt, vorlage, zeile + 1)
    _schriftfeld(blatt, projekt, leiste_nr, vorlage, zeug, fuss, len(spalten))
    blatt.freeze_panes = blatt.cell(row=erste_zeile, column=1)
    _seiteneinrichtung(blatt)
    return blatt


def _kabelbloecke(
    blatt, projekt: Projekt, leiste_nr: int, vorlage: Blattvorlage, zeug, erste: int
) -> None:
    """Die Kabelklammer: ueber ihre Zeilen verbunden, wie auf dem Blatt."""
    stelle = next(
        (i for i, s in enumerate(vorlage.spalten, 1) if s.feld == "kabel"), None
    )
    if stelle is None:
        return

    for kabel in projekt.verteiler.kabel:
        for abschnitt in kabel.abschnitte:
            if abschnitt.leiste != leiste_nr:
                continue
            von = erste + (abschnitt.start_da - 1) * 2
            bis = von + abschnitt.anzahl_da * 2 - 1
            beschriftung = kabel.beschriftung
            if kabel.gekuerzt:
                beschriftung += (
                    f" (gekürzt: {kabel.da_aufgelegt} von {kabel.da_gesamt} DA)"
                )
            blatt.merge_cells(
                start_row=von, start_column=stelle, end_row=bis, end_column=stelle
            )
            feld = blatt.cell(row=von, column=stelle, value=beschriftung)
            feld.font = zeug.schrift(vorlage.schrift_klein)
            feld.alignment = zeug.lage(
                "mitte" if vorlage.kabel_gedreht else "links",
                gedreht=90 if vorlage.kabel_gedreht else 0,
            )
            feld.fill = zeug.Fuellung("solid", start_color=_hex(FUELLUNG_KABEL))
            for zeile in range(von, bis + 1):
                blatt.cell(row=zeile, column=stelle).border = zeug.rahmen(
                    left=zeug.strich(LINIE), right=zeug.strich(LINIE),
                    top=zeug.strich(LINIE) if zeile == von else None,
                    bottom=zeug.strich(LINIE) if zeile == bis else None,
                )


def _spaltenteilung(
    blatt, zeug, anteile: list[float], von: int = 1, bis: int = 0,
    mindestens: int = 1,
) -> list[int]:
    """Wo die Bloecke des Schriftfeldes enden - in Spaltennummern.

    Gemessen wird an den Breiten der Tabelle darueber, damit der
    Trennstrich ungefaehr dort steht, wo ihn das gezeichnete Blatt hat.
    Jeder Block bekommt mindestens so viele Spalten, wie er braucht -
    eine fuer die Beschriftung, eine fuer den Wert; reicht die Tabelle
    dafuer nicht, bekommt jeder eine.
    """
    anzahl = len(anteile)
    spaltenzahl = bis - von + 1
    if spaltenzahl < anzahl * mindestens:
        mindestens = 1
    breiten = [
        blatt.column_dimensions[zeug.spaltenname(i)].width or 8.43
        for i in range(von, bis + 1)
    ]
    gesamt = sum(breiten) or 1.0

    grenzen: list[int] = []
    ziel = 0.0
    for anteil in anteile[:-1]:
        ziel += anteil * gesamt
        gelaufen = 0.0
        stelle = 1
        for k, breite in enumerate(breiten, 1):
            if gelaufen + breite / 2 >= ziel:
                break
            gelaufen += breite
            stelle = k
        # Kein Block darf leer ausgehen, und fuer die folgenden muss
        # genug uebrig bleiben.
        untergrenze = (len(grenzen) + 1) * mindestens
        obergrenze = spaltenzahl - (anzahl - len(grenzen) - 1) * mindestens
        grenzen.append(von - 1 + max(untergrenze, min(stelle, obergrenze)))
    grenzen.append(bis)
    return grenzen


def _block_rahmen(blatt, zeug, oben: int, unten: int, von: int, bis: int) -> None:
    """Zieht einen Rahmen um einen Bereich - Zelle fuer Zelle."""
    for zeile in range(oben, unten + 1):
        for spalte in range(von, bis + 1):
            feld = blatt.cell(row=zeile, column=spalte)
            vorher = feld.border
            feld.border = zeug.rahmen(
                left=zeug.strich(LINIE) if spalte == von else vorher.left,
                right=zeug.strich(LINIE) if spalte == bis else vorher.right,
                top=zeug.strich(LINIE) if zeile == oben else vorher.top,
                bottom=zeug.strich(LINIE) if zeile == unten else vorher.bottom,
            )


def _querstrich(blatt, zeug, zeile: int, von: int, bis: int) -> None:
    """Ein Querstrich im Schriftfeld - in Excel der untere Rand der Zeile."""
    for spalte in range(von, bis + 1):
        feld = blatt.cell(row=zeile, column=spalte)
        vorher = feld.border
        feld.border = zeug.rahmen(
            left=vorher.left, right=vorher.right, top=vorher.top,
            bottom=zeug.strich(LINIE),
        )


def _beschriftet(blatt, zeug, zeile: int, spalte: int, text: str,
                 vorlage: Blattvorlage, ausrichtung: str = "links") -> None:
    """Eine Beschriftung des Schriftfeldes - klein und leise wie am Blatt."""
    feld = blatt.cell(row=zeile, column=spalte, value=text)
    feld.font = zeug.schrift(vorlage.schrift_klein, False, "#555555")
    feld.alignment = zeug.lage(ausrichtung)


def _wertfeld(blatt, zeug, zeile: int, spalte: int, text: str,
              vorlage: Blattvorlage, ausrichtung: str = "links") -> None:
    feld = blatt.cell(row=zeile, column=spalte, value=text)
    feld.font = zeug.schrift(vorlage.schrift_normal)
    feld.alignment = zeug.lage(ausrichtung)


def _schriftfeld(
    blatt, projekt: Projekt, leiste_nr: int | None, vorlage: Blattvorlage, zeug,
    zeile: int, spaltenzahl: int, blattart: str = "leisten",
) -> int:
    """Das Schriftfeld unter der Tabelle - in der Bauart der Blattvorlage.

    Excel kann keine freie Zeichnung, aber einen gerahmten Block aus
    Zellen: dieselben Felder an denselben Plaetzen wie auf dem Blatt.
    Zurueck kommt die Zeile darunter.
    """
    if vorlage.schriftfeld_art == "beispiel":
        return _schriftfeld_beispiel(
            blatt, projekt, leiste_nr, vorlage, zeug, zeile, spaltenzahl, blattart
        )
    return _schriftfeld_spalten(
        blatt, projekt, leiste_nr, vorlage, zeug, zeile, spaltenzahl, blattart
    )


def _schriftfeld_spalten(
    blatt, projekt: Projekt, leiste_nr: int | None, vorlage: Blattvorlage, zeug,
    zeile: int, spaltenzahl: int, blattart: str,
) -> int:
    """Die Bloecke nebeneinander, jeder mit Beschriftung und Wert."""
    bloecke = [
        block for block in schriftfeldspalten(
            projekt, leiste_nr, vorlage, blattart,
            _blatttitel(projekt, leiste_nr, blattart),
        ) if block
    ]
    if not bloecke:
        return zeile
    grenzen = _spaltenteilung(
        blatt, zeug,
        schriftfeldanteile(len(bloecke), vorlage.schriftfeld_anteile),
        bis=spaltenzahl, mindestens=2,
    )
    hoehe = max(len(block) for block in bloecke)
    von = 1
    for block, bis in zip(bloecke, grenzen):
        for i, (name, wert, ausrichtung) in enumerate(block):
            if name == STRICH:
                _querstrich(blatt, zeug, zeile + i, von, bis)
                continue
            if bis == von:
                # Nur eine Spalte: Beschriftung und Wert in einer Zelle.
                _wertfeld(
                    blatt, zeug, zeile + i, von,
                    f"{name}: {wert}".strip(" :") if name else str(wert), vorlage,
                    ausrichtung,
                )
                continue
            _beschriftet(blatt, zeug, zeile + i, von, name, vorlage)
            _wertfeld(
                blatt, zeug, zeile + i, von + 1, str(wert), vorlage, ausrichtung
            )
            if bis > von + 1:
                blatt.merge_cells(
                    start_row=zeile + i, start_column=von + 1,
                    end_row=zeile + i, end_column=bis,
                )
        _block_rahmen(blatt, zeug, zeile, zeile + hoehe - 1, von, bis)
        von = bis + 1
    for i in range(hoehe):
        blatt.row_dimensions[zeile + i].height = zeilenhoehe(
            vorlage.schriftfeld_hoehe / max(hoehe, 1)
        )
    return zeile + hoehe


def _schriftfeld_beispiel(
    blatt, projekt: Projekt, leiste_nr: int | None, vorlage: Blattvorlage, zeug,
    zeile: int, spaltenzahl: int, blattart: str,
) -> int:
    """Das Schriftfeld der Mustertabellen, im Raster nachgebaut.

    Links oben die Matrix aus Datum und Name, links unten der Verteiler,
    rechts zwei Kaesten mit den Anschriften - dieselbe Aufteilung wie auf
    dem gezeichneten Blatt.
    """
    matrix, verteiler, anschriften = beispielschriftfeld(
        projekt, leiste_nr, blattart, vorlage
    )
    links = _spaltenteilung(
        blatt, zeug, schriftfeldanteile(2, vorlage.schriftfeld_anteile),
        bis=spaltenzahl, mindestens=2,
    )[0]
    # Die linke Haelfte traegt drei Felder: Beschriftung, Datum, Name -
    # in denselben Anteilen wie auf dem Blatt.
    rest = (1.0 - BESCHRIFTUNGSANTEIL) / 2
    label_bis, datum_bis, name_bis = _spaltenteilung(
        blatt, zeug, [BESCHRIFTUNGSANTEIL, rest, rest], bis=links,
    )
    datum = min(label_bis + 1, links)
    name = min(datum_bis + 1, links)

    oben = zeile
    hoehe = max(len(matrix) + 1, len(verteiler), 3)
    mitte = oben + hoehe          # erste Zeile der unteren Haelfte
    unten = mitte + hoehe - 1

    # Oben links: die Kopfzeile aus Datum und Name, darunter die Zeilen.
    for spalte, letzte, ueberschrift in (
        (datum, datum_bis, "Datum"), (name, name_bis, "Name"),
    ):
        if spalte <= links:
            _beschriftet(blatt, zeug, oben, spalte, ueberschrift, vorlage, "mitte")
            if letzte > spalte:
                blatt.merge_cells(
                    start_row=oben, start_column=spalte,
                    end_row=oben, end_column=letzte,
                )
    for i, (beschriftung, wert_datum, wert_name) in enumerate(matrix, 1):
        _beschriftet(blatt, zeug, oben + i, 1, beschriftung, vorlage)
        _wertfeld(blatt, zeug, oben + i, datum, wert_datum, vorlage, "mitte")
        if name <= links:
            _wertfeld(blatt, zeug, oben + i, name, wert_name, vorlage, "mitte")
        for erste, letzte in ((datum, datum_bis), (name, name_bis)):
            if letzte > erste:
                blatt.merge_cells(
                    start_row=oben + i, start_column=erste,
                    end_row=oben + i, end_column=letzte,
                )

    # Unten links: Verteiler, Nummer, Leiste.
    for i, (beschriftung, wert) in enumerate(verteiler):
        _beschriftet(blatt, zeug, mitte + i, 1, beschriftung, vorlage)
        _wertfeld(blatt, zeug, mitte + i, datum, str(wert), vorlage, "mitte")
        if links > datum:
            blatt.merge_cells(
                start_row=mitte + i, start_column=datum,
                end_row=mitte + i, end_column=links,
            )

    # Rechts: zwei Kaesten uebereinander, die Anschrift zeilenweise.
    if links < spaltenzahl:
        for i, (beschriftung, zeilen) in enumerate(anschriften):
            erste = oben if i == 0 else mitte
            blatt.merge_cells(
                start_row=erste, start_column=links + 1,
                end_row=erste + hoehe - 1, end_column=spaltenzahl,
            )
            feld = blatt.cell(row=erste, column=links + 1)
            if zeilen:
                feld.value = "\n".join(zeilen)
                feld.font = zeug.schrift(vorlage.schrift_normal)
                feld.alignment = zeug.Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            else:
                feld.value = beschriftung
                feld.font = zeug.schrift(vorlage.schrift_klein, False, "#555555")
                feld.alignment = zeug.lage("links")
            _block_rahmen(
                blatt, zeug, erste, erste + hoehe - 1, links + 1, spaltenzahl
            )

    _block_rahmen(blatt, zeug, oben, unten, 1, links)
    _block_rahmen(blatt, zeug, mitte, unten, 1, links)
    for i in range(oben, unten + 1):
        blatt.row_dimensions[i].height = zeilenhoehe(
            vorlage.schriftfeld_hoehe / (2 * hoehe)
        )
    return unten + 1


def _bis_zum_blattfuss(
    blatt, vorlage: Blattvorlage, zeile: int, quer: bool = False
) -> int:
    """Schiebt das Schriftfeld an den Fuss der Seite.

    Auf dem gezeichneten Blatt steht es unten am Rand, und der freie
    Platz liegt darueber. Excel kennt keine feste Lage - also wird der
    Rest als Luft eingelegt: eine Zeile, so hoch wie noch fehlt, und wenn
    das mehr ist, als eine Zeile tragen darf, mehrere.

    Zurueck kommt die Zeile, in der das Schriftfeld beginnt.
    """
    hoch = 297.0 if not quer else 210.0
    nutzbar = hoch - 2 * SEITENRAND_MM
    genutzt = sum(
        (blatt.row_dimensions[z].height or VORGABEHOEHE_PT)
        for z in range(1, zeile)
    ) / PT_JE_MM
    rest = nutzbar - genutzt - vorlage.schriftfeld_hoehe
    if rest <= 1.0:
        return zeile
    zeilen = max(1, int(rest * PT_JE_MM / HOECHSTE_ZEILE_PT) + 1)
    for i in range(zeilen):
        blatt.row_dimensions[zeile + i].height = zeilenhoehe(rest / zeilen)
    return zeile + zeilen


def _blatttitel(projekt: Projekt, leiste_nr: int | None, blattart: str) -> str:
    """Wie das Blatt heisst - so wie das gezeichnete auch."""
    if blattart == "belegung":
        return "Verteilerbelegung"
    if leiste_nr is None:
        return ""
    return f"Leiste {projekt.leistenname(leiste_nr)}"
