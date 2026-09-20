"""Tests zur Excel-Mappe im Blattbild.

Geprueft wird, was Excel von der Blattvorlage uebernimmt: Spaltenfolge und
-breite, Schriftgroessen, der verbundene Kabelblock samt Drehung und die
Seiteneinrichtung. Wie es dann aussieht, entscheidet Excel selbst - das
laesst sich hier nicht nachmessen.
"""

import pathlib
import tempfile
import unittest

from openpyxl import load_workbook

from verteilerly.ausgabe import xlsx_schreiben
from verteilerly.ausgabe.mappe import spaltenbreite, zeilenhoehe
from verteilerly.blattvorlagen import STANDARD, Blattvorlagen
from verteilerly.modell import Aderadresse, Projekt, Verteiler
from verteilerly.pfade import mitgeliefert
from . import BEISPIELVORLAGE


def _projekt() -> Projekt:
    v = Verteiler.aus_geometrie(1, 2, 10, name="Hauptverteiler", nummer="V-01")
    v.kabel_anlegen("Linie 1 — Melderkreis", 1, 1, 5, kabeltyp="jystyv")
    v.kabel_anlegen("Zuleitung BMZ", 1, 7, 3, da_gesamt=4)
    v.rangierung_anlegen(Aderadresse(1, 1, "a"), Aderadresse(2, 2, "a"))
    v.ader(Aderadresse(1, 1, "a")).bezeichnung = "Melder 1a"
    return Projekt(verteiler=v, briefkopf={"firma": "Musterfirma"})


class TestQuerstrich(unittest.TestCase):
    """In Excel ist der Querstrich der untere Rand einer Zeile."""

    def test_aus_der_marke_wird_ein_rand(self):
        import dataclasses

        vorlage = dataclasses.replace(
            STANDARD, schriftfeld=(("firma", "strich", "verteiler"),)
        )
        with tempfile.TemporaryDirectory() as ordner:
            ziel = pathlib.Path(ordner) / "probe.xlsx"
            xlsx_schreiben(_projekt(), ziel, vorlage)
            blatt = load_workbook(ziel)["Leiste 1"]
            texte = [
                zelle.value for zeile in blatt.iter_rows() for zelle in zeile
                if isinstance(zelle.value, str)
            ]
            self.assertFalse([t for t in texte if "\x00" in t])
            mit_rand = [
                zelle for zeile in blatt.iter_rows() for zelle in zeile
                if zelle.border.bottom and zelle.border.bottom.style
            ]
            self.assertTrue(mit_rand, "keine Zeile mit unterem Rand")


class TestMasse(unittest.TestCase):
    def test_millimeter_werden_umgerechnet(self):
        # 25.4 mm sind ein Zoll: 96 Bildpunkte, 72 Punkt.
        self.assertAlmostEqual(zeilenhoehe(25.4), 72.0, places=1)
        self.assertAlmostEqual(spaltenbreite(25.4), (96 - 5) / 7, places=1)

    def test_keine_spalte_verschwindet(self):
        self.assertGreaterEqual(spaltenbreite(0.1), 1.0)


class TestBlattbild(unittest.TestCase):
    def setUp(self):
        self.ordner = tempfile.TemporaryDirectory()
        self.addCleanup(self.ordner.cleanup)
        self.ziel = pathlib.Path(self.ordner.name) / "probe.xlsx"
        self.projekt = _projekt()

    def _mappe(self, vorlage=STANDARD):
        xlsx_schreiben(self.projekt, self.ziel, vorlage)
        return load_workbook(self.ziel)

    def test_je_leiste_ein_blatt(self):
        mappe = self._mappe()
        self.assertIn("Leiste 1", mappe.sheetnames)
        self.assertIn("Leiste 2", mappe.sheetnames)
        self.assertEqual(mappe.sheetnames[0], "Deckblatt")

    def test_spalten_folgen_der_blattvorlage(self):
        blatt = self._mappe()["Leiste 1"]
        self.assertEqual(
            [zelle.value or "" for zelle in blatt[1]],
            [spalte.ueberschrift for spalte in STANDARD.spalten],
        )

    def test_stiftueberschrift_reicht_ueber_beide_spalten(self):
        blatt = self._mappe()["Leiste 1"]
        self.assertIn("A1:B1", [str(r) for r in blatt.merged_cells.ranges])
        self.assertEqual(blatt["A1"].value, "Stift")

    def test_zwei_zeilen_je_stift(self):
        blatt = self._mappe()["Leiste 1"]
        self.assertEqual(blatt.cell(row=2, column=1).value, "1")
        self.assertEqual(blatt.cell(row=2, column=2).value, "a")
        self.assertFalse(blatt.cell(row=3, column=1).value)   # b-Zeile
        self.assertEqual(blatt.cell(row=3, column=2).value, "b")
        self.assertEqual(blatt.cell(row=2, column=4).value, "Melder 1a")
        self.assertEqual(blatt.cell(row=2, column=5).value, "2 2a")

    def test_kabelblock_ist_verbunden_und_gedreht(self):
        blatt = self._mappe()["Leiste 1"]
        bereiche = [str(r) for r in blatt.merged_cells.ranges]
        self.assertIn("G2:G11", bereiche)     # 5 DA sind zehn Zeilen
        zelle = blatt["G2"]
        self.assertEqual(zelle.value, "Linie 1 — Melderkreis")
        self.assertEqual(zelle.alignment.text_rotation, 90)

    def test_gekuerztes_kabel_sagt_es(self):
        blatt = self._mappe()["Leiste 1"]
        self.assertIn("gekürzt: 3 von 4 DA", blatt["G14"].value)

    def test_waagerechte_vorlage_dreht_nicht(self):
        knapp = Blattvorlagen.laden(mitgeliefert() / "blattvorlagen").holen("knapp")
        blatt = self._mappe(knapp)["Leiste 1"]
        stelle = [s.feld for s in knapp.spalten].index("kabel") + 1
        zelle = blatt.cell(row=2, column=stelle)
        self.assertEqual(zelle.alignment.text_rotation, 0)
        self.assertEqual(zelle.value, "Linie 1 — Melderkreis")

    def test_spaltenbreiten_kommen_aus_der_vorlage(self):
        blatt = self._mappe()["Leiste 1"]
        breiten = STANDARD.breiten(210.0 - 2 * STANDARD.rand)
        self.assertAlmostEqual(
            blatt.column_dimensions["A"].width, spaltenbreite(breiten[0]), places=1
        )
        # Die Kabelspalte ist die schmalste - so steht es in der Vorlage.
        self.assertLess(
            blatt.column_dimensions["G"].width, blatt.column_dimensions["D"].width
        )

    def test_schriftgroessen_kommen_aus_der_vorlage(self):
        """Die Tabelle beginnt in Zeile 1 - kein Blattkopf darüber."""
        blatt = self._mappe()["Leiste 1"]
        self.assertEqual(blatt["A1"].value, "Stift")
        self.assertEqual(blatt["A1"].font.sz, STANDARD.schrift_klein)   # Kopfzeile
        self.assertEqual(blatt["D2"].font.sz, STANDARD.schrift_normal)  # Aderzeile
        self.assertTrue(blatt["A1"].font.b)

    @staticmethod
    def _schriftfeldbeginn(blatt) -> int:
        """Die erste Zeile unter der Tabelle - dort steht das Schriftfeld."""
        letzte = max(
            (z.row for zeile in blatt.iter_rows(min_col=2, max_col=2)
             for z in zeile if z.value in ("a", "b")),
            default=1,
        )
        return letzte + 1

    def _schriftfeldpaare(self, blatt, ab: int = 0) -> dict:
        """Beschriftung und Wert daneben - gleich in welchem Block."""
        ab = ab or self._schriftfeldbeginn(blatt)
        paare = {}
        for zeile in range(ab, ab + 20):
            for spalte in range(1, 9):
                name = blatt.cell(row=zeile, column=spalte).value
                wert = blatt.cell(row=zeile, column=spalte + 1).value
                if name:
                    paare[str(name)] = wert
        return paare

    def test_schriftfeld_steht_unter_der_tabelle(self):
        """Die Blöcke stehen nebeneinander wie auf dem Blatt."""
        blatt = self._mappe()["Leiste 1"]
        gefunden = self._schriftfeldpaare(blatt)
        self.assertEqual(gefunden.get("Firma"), "Musterfirma")
        self.assertEqual(gefunden.get("Verteiler"), "Hauptverteiler")
        self.assertEqual(gefunden.get("Leiste"), "1")

    def test_das_schriftfeld_traegt_einen_rahmen(self):
        """Auf dem Blatt ist es ein Kasten — im Raster sind es Ränder."""
        blatt = self._mappe()["Leiste 1"]
        ecke = None
        for zeile in range(self._schriftfeldbeginn(blatt), blatt.max_row + 1):
            if blatt.cell(row=zeile, column=1).value:
                ecke = blatt.cell(row=zeile, column=1)
                break
        self.assertIsNotNone(ecke, "kein Schriftfeld gefunden")
        self.assertTrue(ecke.border.left.style, "kein linker Rand")
        self.assertTrue(ecke.border.top.style, "kein oberer Rand")

    def test_seite_ist_auf_a4_eingerichtet(self):
        mappe = self._mappe()
        blatt = mappe["Leiste 1"]
        self.assertEqual(blatt.page_setup.orientation, "portrait")
        self.assertEqual(str(blatt.page_setup.paperSize), str(blatt.PAPERSIZE_A4))
        self.assertEqual(blatt.page_setup.fitToWidth, 1)
        # Die Verteilerbelegung folgt dem gezeichneten Blatt: eine Bucht
        # hochkant, ab zwei quer.
        self.assertEqual(
            mappe["Verteilerbelegung"].page_setup.orientation, "portrait"
        )

    @staticmethod
    def _leistenzeile(blatt, nummer: int = 1) -> int:
        """In welcher Zeile die Leiste steht - die Luft dazwischen zählt mit."""
        for zeile in blatt.iter_rows(min_col=1, max_col=1):
            if zeile[0].value == f"Leiste {nummer}":
                return zeile[0].row
        raise AssertionError(f"Leiste {nummer} steht nicht im Blatt")

    def test_belegung_zeigt_die_kabel_als_block(self):
        blatt = self._mappe()["Verteilerbelegung"]
        # Überschrift oben, darunter die Bucht mit ihren Leisten.
        self.assertEqual(blatt.cell(row=1, column=1).value, "Verteilerbelegung")
        texte = [z.value for zeile in blatt.iter_rows(max_col=1) for z in zeile]
        self.assertIn("Bucht 1", texte)
        zeile = self._leistenzeile(blatt)
        self.assertIn(
            f"B{zeile}:F{zeile}", [str(r) for r in blatt.merged_cells.ranges]
        )
        self.assertEqual(blatt[f"B{zeile}"].value, "Linie 1 — Melderkreis")

    def test_der_kabelname_steht_mittig_und_bricht_um(self):
        """Wie auf dem Blatt - nicht klein oben links."""
        from verteilerly.layout import (
            belegungsmasse,
            kabelfeld,
            umgebrochene_zeilen,
        )

        blatt = self._mappe()["Verteilerbelegung"]
        zeile = self._leistenzeile(blatt)
        feld = blatt[f"B{zeile}"]
        masse = belegungsmasse(self.projekt.verteiler.buchten, STANDARD)
        self.assertEqual(feld.alignment.horizontal, "center")
        self.assertEqual(feld.alignment.vertical, "center")
        self.assertTrue(feld.alignment.wrap_text)
        # Fuenf DA breit: dieselbe Rechnung wie fuer Vorschau und PDF.
        platz = kabelfeld(
            0.0, 0.0, 5 * masse.zelle_breite, masse.zelle_hoehe, STANDARD
        )
        groesse, zeilen = umgebrochene_zeilen(
            "Linie 1 — Melderkreis", platz.breite, platz.hoehe,
            platz.groesse, platz.einzeilig_ab,
        )
        self.assertEqual(feld.value, "\n".join(zeilen))
        self.assertEqual(feld.font.size, groesse)

    def test_ein_schmaler_block_verkleinert_die_schrift(self):
        """Excel kennt die Breite des verbundenen Blocks nicht.

        Frueher stand der Name unverkleinert darin und lief darueber
        hinaus - die Mappe rechnet die Groesse darum selbst aus.
        """
        from verteilerly.layout import belegungsmasse, kabelfeld

        v = self.projekt.verteiler
        v.kabel_anlegen("Melderkreis Erdgeschoss West", 2, 1, 1)
        blatt = self._mappe()["Verteilerbelegung"]
        masse = belegungsmasse(v.buchten, STANDARD)
        gross = kabelfeld(0.0, 0.0, masse.zelle_breite, masse.zelle_hoehe).groesse

        feld = None
        for zeile in blatt.iter_rows():
            for zelle in zeile:
                if zelle.value and str(zelle.value).startswith("Melderkreis"):
                    feld = zelle
        self.assertIsNotNone(feld, "der Kabelname steht nicht im Blatt")
        self.assertLess(feld.font.size, gross)
        self.assertIn("\n", feld.value)

    def test_die_leisten_fuellen_das_blatt_wie_im_pdf(self):
        """Die Zellenmaße kommen aus dem Blattmodell."""
        from verteilerly.ausgabe.mappe import zeilenhoehe
        from verteilerly.layout import belegungsmasse

        blatt = self._mappe()["Verteilerbelegung"]
        masse = belegungsmasse(self.projekt.verteiler.buchten, STANDARD)
        zeile = self._leistenzeile(blatt)
        self.assertAlmostEqual(
            blatt.row_dimensions[zeile].height,
            zeilenhoehe(masse.zelle_hoehe), places=1,
        )

    def test_anschrift_steht_zeilenweise_im_schriftfeld(self):
        """Firma und Kunde wie in den Mustertabellen."""
        self.projekt.briefkopf["firma"] = (
            "Firma Beliebig\nBeliebige Straße 7\n77777 Beliebiger Ort"
        )
        blatt = self._mappe()["Leiste 1"]
        gefunden = [
            (blatt.cell(row=z, column=1).value, blatt.cell(row=z, column=2).value)
            for z in range(self._schriftfeldbeginn(blatt), blatt.max_row + 1)
        ]
        self.assertIn(("Firma", "Firma Beliebig"), gefunden)
        self.assertIn((None, "Beliebige Straße 7"), gefunden)
        self.assertIn((None, "77777 Beliebiger Ort"), gefunden)

    def test_anschrift_bricht_in_den_projektdaten_um(self):
        self.projekt.briefkopf["firma"] = "Firma Beliebig\nBeliebige Straße 7"
        blatt = self._mappe()["Projektdaten"]
        zellen = [
            blatt.cell(row=z, column=2)
            for z in range(1, blatt.max_row + 1)
            if isinstance(blatt.cell(row=z, column=2).value, str)
            and "\n" in blatt.cell(row=z, column=2).value
        ]
        self.assertTrue(zellen)
        self.assertTrue(zellen[0].alignment.wrap_text)

    def test_daten_stehen_weiter_dahinter(self):
        """Das Blattbild tritt nicht an die Stelle der Daten."""
        mappe = self._mappe()
        self.assertEqual(mappe["Daten je Ader"].max_row, 41)   # 20 DA, zwei Adern
        self.assertEqual(mappe["Daten je Ader"].cell(row=2, column=5).value, "1 1a")


if __name__ == "__main__":
    unittest.main()


class TestSchriftfeldWieAufDemBlatt(unittest.TestCase):
    """Dieselben Felder an denselben Plätzen wie im PDF.

    Bis hierher war das Schriftfeld der Mappe eine schlichte Liste aus
    Beschriftung und Wert - die Bauart "beispiel" baut es aber aus einer
    Matrix, dem Verteiler und zwei Anschriften.
    """

    @classmethod
    def setUpClass(cls):
        cls.vorlage = BEISPIELVORLAGE

    def setUp(self):
        self.ordner = tempfile.TemporaryDirectory()
        self.addCleanup(self.ordner.cleanup)
        ziel = pathlib.Path(self.ordner.name) / "probe.xlsx"
        self.projekt = Projekt(
            verteiler=Verteiler.aus_geometrie(
                1, 2, 10, name="Hauptverteiler", nummer="1.1"
            ),
            briefkopf={
                "firma": "Netzbau GmbH\nIndustriestraße 12\n12345 Musterstadt",
                "kunde": "Firma Meier\nHauptstraße 3",
                "erstellt_am": "01.02.2026", "erstellt_von": "T. Mirik",
                "geprueft_am": "03.02.2026", "geprueft_von": "K. Prüfer",
            },
        )
        xlsx_schreiben(self.projekt, ziel, self.vorlage)
        self.mappe = load_workbook(ziel)

    def _felder(self, blatt) -> dict:
        """Jede Beschriftung mit dem Wert, der rechts daneben steht."""
        paare = {}
        for zeile in blatt.iter_rows():
            for z in zeile:
                if not z.value:
                    continue
                for weiter in range(z.column + 1, blatt.max_column + 1):
                    wert = blatt.cell(row=z.row, column=weiter).value
                    if wert not in (None, ""):
                        paare.setdefault(str(z.value), wert)
                        break
        return paare

    def test_die_matrix_steht_oben_links(self):
        blatt = self.mappe["Leiste 1"]
        felder = self._felder(blatt)
        self.assertEqual(felder.get("Erstellt"), "01.02.2026")
        self.assertEqual(felder.get("Geprüft"), "03.02.2026")
        # Die Kopfzeile der Matrix: Datum und Name über den Spalten.
        texte = {z.value for zeile in blatt.iter_rows() for z in zeile if z.value}
        self.assertIn("Datum", texte)
        self.assertIn("Name", texte)
        self.assertIn("T. Mirik", texte)

    def test_der_verteiler_steht_unten_links(self):
        felder = self._felder(self.mappe["Leiste 1"])
        self.assertEqual(felder.get("Verteiler:"), "Hauptverteiler")
        self.assertEqual(felder.get("Verteiler-Nr."), "1.1")
        self.assertEqual(felder.get("Leiste:"), "1")

    def test_die_anschriften_stehen_rechts_in_zwei_kaesten(self):
        blatt = self.mappe["Leiste 1"]
        kaesten = [
            z for zeile in blatt.iter_rows() for z in zeile
            if isinstance(z.value, str) and "\n" in z.value
        ]
        self.assertEqual(len(kaesten), 2, "nicht zwei Anschriftenkästen")
        self.assertTrue(kaesten[0].value.startswith("Netzbau GmbH"))
        self.assertTrue(kaesten[1].value.startswith("Firma Meier"))
        for kasten in kaesten:
            with self.subTest(zelle=kasten.coordinate):
                # Mehrzeilig heißt in Excel: umbrechen und verbunden.
                self.assertTrue(kasten.alignment.wrap_text)
                self.assertTrue(
                    any(kasten.coordinate == str(m).split(":")[0]
                        for m in blatt.merged_cells.ranges),
                    "der Kasten ist nicht verbunden",
                )

    def test_das_schriftfeld_ist_ein_kasten(self):
        blatt = self.mappe["Leiste 1"]
        ecke = next(
            z for zeile in blatt.iter_rows() for z in zeile
            if z.value == "Verteiler:"
        )
        self.assertTrue(ecke.border.left.style)

    def test_auch_die_verteilerbelegung_traegt_es(self):
        """Im PDF steht es dort ebenfalls am Blattfuß."""
        felder = self._felder(self.mappe["Verteilerbelegung"])
        self.assertEqual(felder.get("Verteiler:"), "Hauptverteiler")
        self.assertEqual(felder.get("Erstellt"), "01.02.2026")
        # Dieses Blatt zeigt alle Leisten, also keine einzelne.
        self.assertNotIn("Leiste:", felder)

    @staticmethod
    def _leistenzeile(blatt, nummer: int = 1) -> int:
        for zeile in blatt.iter_rows(min_col=1, max_col=1):
            if zeile[0].value == f"Leiste {nummer}":
                return zeile[0].row
        raise AssertionError(f"Leiste {nummer} steht nicht im Blatt")

    def test_keine_da_nummern_wenn_die_vorlage_keine_fuehrt(self):
        """Die Vorlage sagt „da_nummern: false" - dann steht dort nichts."""
        blatt = self.mappe["Verteilerbelegung"]
        self.assertFalse(self.vorlage.da_nummern)
        zeile = self._leistenzeile(blatt)
        zellen = [blatt.cell(row=zeile, column=s).value for s in range(2, 12)]
        self.assertEqual([z for z in zellen if z], [])
        # Der Kasten um die Leiste bleibt trotzdem.
        self.assertTrue(blatt.cell(row=zeile, column=2).border.top.style)

    def test_mit_da_nummern_stehen_sie_darin(self):
        from dataclasses import replace

        ziel = pathlib.Path(self.ordner.name) / "nummern.xlsx"
        xlsx_schreiben(self.projekt, ziel, replace(self.vorlage, da_nummern=True))
        blatt = load_workbook(ziel)["Verteilerbelegung"]
        zeile = self._leistenzeile(blatt)
        self.assertEqual(blatt.cell(row=zeile, column=2).value, "1")

    def test_die_zeilenhoehe_kommt_vom_blatt(self):
        """Die Tabelle nutzt die Seite aus wie im PDF."""
        from verteilerly.ausgabe.mappe import zeilenhoehe
        from verteilerly.layout import zeilenhoehe_leiste

        blatt = self.mappe["Leiste 1"]
        soll = zeilenhoehe(zeilenhoehe_leiste(10, self.vorlage))
        self.assertAlmostEqual(blatt.row_dimensions[2].height, soll, places=1)
        # Und das füllt die Seite: 20 Zeilen tragen mehr als die Hälfte.
        gesamt = sum(
            (blatt.row_dimensions[z].height or 15)
            for z in range(1, blatt.max_row + 1)
        )
        self.assertGreater(gesamt / (72 / 25.4), 0.7 * (297 - 25.4))

    def test_die_buchten_stehen_nebeneinander(self):
        """Wie im PDF: zwei Buchten teilen sich ein Blatt quer."""
        ziel = pathlib.Path(self.ordner.name) / "zwei.xlsx"
        projekt = Projekt(
            verteiler=Verteiler.aus_geometrie(2, 3, 10, name="Zwei"),
            briefkopf={"firma": "Netzbau GmbH"},
        )
        xlsx_schreiben(projekt, ziel, self.vorlage)
        blatt = load_workbook(ziel)["Verteilerbelegung"]
        self.assertEqual(blatt.page_setup.orientation, "landscape")
        zeilen = {
            z.value: z.row
            for zeile in blatt.iter_rows(max_col=15) for z in zeile
            if isinstance(z.value, str) and z.value.startswith("Leiste ")
        }
        # Leiste 1 (Bucht 1) und Leiste 4 (Bucht 2) stehen auf gleicher Höhe.
        self.assertEqual(zeilen["Leiste 1"], zeilen["Leiste 4"])
        spalten = {
            z.value: z.column
            for zeile in blatt.iter_rows(max_col=30) for z in zeile
            if isinstance(z.value, str) and z.value.startswith("Leiste ")
        }
        self.assertGreater(spalten["Leiste 4"], spalten["Leiste 1"] + 10)

    def test_je_blatt_der_zeichnung_eines_in_der_mappe(self):
        """Ab drei Buchten bekommt jede ihr eigenes Blatt."""
        ziel = pathlib.Path(self.ordner.name) / "drei.xlsx"
        projekt = Projekt(verteiler=Verteiler.aus_geometrie(3, 3, 10, name="Drei"))
        xlsx_schreiben(projekt, ziel, self.vorlage)
        namen = load_workbook(ziel).sheetnames
        for name in ("Bucht 1", "Bucht 2", "Bucht 3"):
            with self.subTest(blatt=name):
                self.assertIn(name, namen)
        self.assertNotIn("Verteilerbelegung", namen)

    def test_das_schriftfeld_steht_am_blattfuss(self):
        """Wie auf dem Blatt: die freie Fläche liegt darüber, nicht darunter."""
        from verteilerly.ausgabe.mappe import (
            SEITENRAND_MM, VORGABEHOEHE_PT,
        )

        PT = 72 / 25.4
        for name in ("Leiste 1", "Verteilerbelegung"):
            with self.subTest(blatt=name):
                blatt = self.mappe[name]
                quer = blatt.page_setup.orientation == "landscape"
                nutzbar = (210.0 if quer else 297.0) - 2 * SEITENRAND_MM
                gesamt = sum(
                    (blatt.row_dimensions[z].height or VORGABEHOEHE_PT)
                    for z in range(1, blatt.max_row + 1)
                ) / PT
                # Die Seite ist voll: das Schriftfeld endet an ihrem Fuß.
                self.assertGreater(gesamt, nutzbar - 4)
                anfang = next(
                    z.row for zeile in blatt.iter_rows(max_col=1) for z in zeile
                    if z.value == "Verteiler:"
                )
                oben = sum(
                    (blatt.row_dimensions[z].height or VORGABEHOEHE_PT)
                    for z in range(1, anfang)
                ) / PT
                # Es beginnt nicht gleich hinter der Tabelle, sondern unten.
                self.assertGreater(oben, nutzbar - self.vorlage.schriftfeld_hoehe - 6)

    def test_ohne_platz_bleibt_es_hinter_der_tabelle(self):
        """Ein volles Blatt bekommt keine Luft eingelegt."""
        from verteilerly.ausgabe.mappe import _bis_zum_blattfuss

        blatt = self.mappe["Leiste 1"]
        letzte = blatt.max_row
        self.assertEqual(_bis_zum_blattfuss(blatt, self.vorlage, letzte), letzte)

    def test_eine_seite_ist_eine_seite(self):
        """Breite und Höhe werden auf ein Blatt gebracht, nicht nur die Breite."""
        blatt = self.mappe["Leiste 1"]
        self.assertEqual(blatt.page_setup.fitToWidth, 1)
        self.assertEqual(blatt.page_setup.fitToHeight, 1)

    def test_das_deckblatt_fuehrt_dieselben_zeilen_wie_im_pdf(self):
        from verteilerly.layout import deckblattzeilen

        blatt = self.mappe["Deckblatt"]
        erwartet = [name for name, _ in deckblattzeilen(self.projekt, self.vorlage)
                    if name]
        texte = [z.value for zeile in blatt.iter_rows() for z in zeile if z.value]
        for name in erwartet:
            with self.subTest(zeile=name):
                self.assertIn(name, texte)
        self.assertIn("Aufbau", texte)
        self.assertNotIn("Buchten", texte, "die alte Kennzahlenliste steht noch da")


if __name__ == "__main__":
    unittest.main()
